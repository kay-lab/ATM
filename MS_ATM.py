#! /usr/bin/env python

#Mass Spectra Automated Trace Maker = MS_ATM
#Version 1.0
#Released to Github on 4/4/2021.

#This Python script will take .CSV mass spectra files (containing only m/z values and
#relative intensity values) and create Excel plots of the mass spectra data.

#The following changes the working directory to the folder in which the Python executable
#is stored in. Disable this if you are not using the executable.
#os.chdir(os.path.dirname(sys.executable))

#Imports important modules and classes.
import csv
import codecs
import glob
import openpyxl
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import (Paragraph, ParagraphProperties, CharacterProperties, 
									Font)

#Introduces user to the program.
print "Welcome to MS_ATM (Mass Spectra Automated Trace Maker)! This program will"
print "prepare mass spectra plots using .CSV files containing m/z values and"
print "the associated relative intensities."
print ""

#Performs the following loop on each .CSV file within a folder.
for File in glob.iglob("*.[Cc][Ss][Vv]"):
	#Tells the user which file is being plotted.
	print "Now plotting "+File+"..."
	print ""
	
	#Puts the m/z values and relative intensities into separate lists.
	mz = []
	intensity = []
	inFile = codecs.open(File, "rU", "utf-16") #codecs used to convert .CSV properly.
	csvReader = csv.reader(inFile)
	for row in csvReader:
		formattedRow = row[0].split("\t")
		mz.append(float(formattedRow[0]))
		intensity.append(float(formattedRow[1]))
	inFile.close()
		
	#Sets up the output Excel file.
	outFile = openpyxl.Workbook()
	sheet = outFile.active
	sheet["A1"] = "m/z"
	sheet["B1"] = "Relative Intensity (%)"
	
	#Puts the m/z values and relative intensity values into the Excel file.
	sheetCounter = 2
	intCounter = 0
	for value in mz:
		sheet["A"+str(sheetCounter)] = value
		sheet["B"+str(sheetCounter)] = intensity[intCounter]
		sheetCounter+=1
		intCounter+=1
	
	#Plots the data via a scatter plot and sets titles for axes.
	chart = ScatterChart()
	chart.y_axis.title = "Relative Intensity (%)"
	chart.x_axis.title = "m/z"
	
	#Sets the chart title text font and size.
	titleFont = Font(typeface="Times New Roman")
	xtextProp = CharacterProperties(latin=titleFont, sz=1200, i=True)
	ytextProp = CharacterProperties(latin=titleFont, sz=1200)
	chart.x_axis.title.tx.rich.p[0].r.rPr = xtextProp
	chart.y_axis.title.tx.rich.p[0].r.rPr = ytextProp
	
	#Sets the axis text font and size.
	axisFont = Font(typeface="Times New Roman")
	fontProp = CharacterProperties(latin=axisFont, sz=1000)
	finalProp = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=fontProp), 
				endParaRPr=fontProp)])
	chart.x_axis.txPr = finalProp
	chart.y_axis.txPr = finalProp
	
	#Creates a data series within the plot and sets the cell in which to draw the plot.
	xvalues = Reference(sheet, min_col=1, min_row=2, max_row=sheetCounter-1)
	yvalues = Reference(sheet, min_col=2, min_row=2, max_row=sheetCounter-1)
	series = Series(values = yvalues, xvalues = xvalues)
	chart.series.append(series)
	sheet.add_chart(chart, "D5")
	
	#Formats the scatter plot with no gridlines.
	noGridLines = GraphicalProperties(ln=LineProperties(noFill=True))
	chart.x_axis.majorGridlines.spPr = noGridLines
	chart.y_axis.majorGridlines.spPr = noGridLines
	
	#Removes the legend from the scatter plot.
	chart.legend = None
	
	#Sets the min and max values for each axis.
	chart.y_axis.scaling.min = 0
	chart.y_axis.scaling.max = 100
	chart.x_axis.scaling.min = 400
	chart.x_axis.scaling.max = 2000
	
	#Sets the y-axis major units.
	chart.y_axis.majorUnit = 100
	
	#Formats the color and width of the scatter plot line.
	series.graphicalProperties.line.solidFill = "800040"
	series.graphicalProperties.line.width = 100
	
	#Saves the output Excel file.
	FileName = File[:-4] #Removes .CSV from File Name
	outFile.save("MS Plot for "+FileName+".xlsx")
	
#Tells the user that the program is finished.
finished = raw_input("MS_ATM is finished! Press 'enter' to close the program.")