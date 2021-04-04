#! /usr/bin/env python

#Automated Trace Maker = ATM
#Version 1.0
#Released to Github on 4/4/2021.

#This Python script will take .CSV HPLC files (containing only retention times and
#absorbance values) and create both unnormalized and normalized Excel plots.

#The following changes the working directory to the folder in which the Python executable
#is stored in. Disable this if you are not using the executable.
#os.chdir(os.path.dirname(sys.executable))

#Imports important modules and classes.
import csv
import codecs
import glob
import openpyxl
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import (Paragraph, ParagraphProperties, CharacterProperties, 
									Font)
import sys
import re

#This allows for the .CSV files to be entered into ATM via ascending order.
numbers = re.compile(r'(\d+)')
def numericalSort(value):
	"""Function for sorting files in numerical and alphabetical ascending order"""
	parts = numbers.split(value)
	parts[1::2] = map(int, parts[1::2])
	return parts

#Introduces user to the program.
print "Welcome to ATM (Automated Trace Maker)! This program will prepare HPLC traces"
print "by using .CSV files containing the retention times and absorbance values."
print ""

#Enables user to enter the same retention times for all files in a folder.
sameRT = False
sameRTAns = raw_input("Do all .CSV files have the same retention times?")
print ""
if sameRTAns == "" or sameRTAns[0].lower() == "y":
	sameRT = True
	userInputs = False
	while userInputs == False:
		minRT = raw_input("Using only numbers, enter the minimum retention time: ")
		print ""
		maxRT = raw_input("Using only numbers, enter the maximum retention time: ")
		print ""
		try:
			minRT = float(minRT)
			maxRT = float(maxRT)
			print "You have entered the following retention times for all files:"
			print ""
			print "Minimum Retention Time = "+str(minRT)
			print "Maximum Retention Time = "+str(maxRT)
			print ""
			correctIn = raw_input("Are these values correct?")
			print ""
			if correctIn == "" or correctIn[0].lower() == "y":
				if minRT < maxRT:
					userInputs = True
				else:
					print "The minimum retention time cannot be larger than or equal to"
					print "the maximum retention time!"
					print "Please enter the retention times again."
					print ""
		except ValueError:
			print "Only numbers can be entered for the retention times!"
			print "Please enter the retention times again."
			print ""

#Performs the following loop on each .CSV file within a folder.
for File in sorted(glob.iglob("*.[Cc][Ss][Vv]"), key=numericalSort):
    #Tells user which file is currently being prepared for plotting.
    print File+" is being prepared for plotting."
    print ""
    
    #Asks user for minimum and maximum retention times that are desired to be plotted
    #and ensures the inputted values are numerical.
    userInputs = False
    while userInputs == False and sameRT == False:
    	minRT = raw_input("Using only numbers, enter the minimum retention time: ")
    	print ""
    	maxRT = raw_input("Using only numbers, enter the maximum retention time: ")
    	print ""
    	try:
    		minRT = float(minRT)
    		maxRT = float(maxRT)
    		print "You have entered the following retention times for"
    		print File+":"
    		print ""
    		print "Minimum Retention Time = "+str(minRT)
    		print "Maximum Retention Time = "+str(maxRT)
    		print ""
    		correctIn = raw_input("Are these values correct?")
    		print ""
    		if correctIn == "" or correctIn[0].lower() == "y":
    			if minRT < maxRT:
    				userInputs = True
    			else:
    				print "The minimum retention time cannot be larger than or equal to"
    				print "the maximum retention time!"
    				print "Please enter the retention times again."
    				print ""
    	except ValueError:
    		print "Only numbers can be entered for the retention times!"
    		print "Please enter the retention times again."
    		print ""
    
    #Tells user that the current file is being plotted.
    print "Now plotting "+File+"..."
    print ""
    
    #Puts the retention times and absorbances into separate lists, based on user's inputs.
    retentionTimes = []
    absorbances = []
    inFile = codecs.open(File, "rU", "utf-16") #codecs used to convert .CSV properly.
    csvReader = csv.reader(inFile)
    for row in csvReader:
    	formattedRow = row[0].split("\t")
    	if float(formattedRow[0]) >= minRT and float(formattedRow[0]) <= maxRT:
    		retentionTimes.append(float(formattedRow[0]))
    		absorbances.append(float(formattedRow[1]))
    inFile.close()
    
    #Finds lowest absorbance value to eventually normalize all data. Also makes sure that
    #user did not enter incorrect retention time information or have an empty CSV file,
    #which would result in an empty absorbance list.
    try:
    	minAbs = min(absorbances)
    except ValueError:
    	print "ERROR! There is no data to plot!"
    	print "You have either entered an empty CSV file, or you have entered the wrong" 
    	print "retention times for this file."
    	print ""
    	print "Please restart the program with the proper file or retention times."
    	exit = raw_input("Press 'enter' to close the program.")
    	sys.exit()
    
    #Sets the lowest absorbance value as zero and corrects all absorbances based on this
    #zero value.
    balancedAbs = []
    for abs in absorbances:
    	balancedAbs.append(abs-minAbs)
    	
    #Finds the maximum balanced absorbance and normalizes all absorbances with this value.
    maxAbs = max(balancedAbs)
    normAbs = []
    for abs in balancedAbs:
    	normAbs.append(abs/maxAbs*100)
    	
    #Sets up the output Excel file.
    outFile = openpyxl.Workbook()
    sheet = outFile.active
    sheet["A1"] = "Retention Time (Min)"
    sheet["B1"] = "Raw Absorbance (mAU)"
    sheet["C1"] = "Normalized Absorbance"
    
    #Puts the retention times and absorbance values into the Excel file.
    sheetCounter = 2
    absCounter = 0
    for time in retentionTimes:
    	sheet["A"+str(sheetCounter)] = time
    	sheet["B"+str(sheetCounter)] = absorbances[absCounter]
    	sheet["C"+str(sheetCounter)] = normAbs[absCounter]
    	sheetCounter += 1
    	absCounter += 1
    
    #Plots the raw data via a scatter plot and sets titles for axes.
    rawChart = ScatterChart()
    rawChart.y_axis.title = "Absorbance (mAU)"
    rawChart.x_axis.title = "Retention Time (min)"
    
    #Sets the chart title text font and size.
    titleFont = Font(typeface="Times New Roman")
    textProp = CharacterProperties(latin=titleFont, sz=1200)
    rawChart.x_axis.title.tx.rich.p[0].r.rPr = textProp
    rawChart.y_axis.title.tx.rich.p[0].r.rPr = textProp
    
    #Sets the axis text font and size.
    axisFont = Font(typeface="Times New Roman")
    fontProp = CharacterProperties(latin=axisFont, sz=1000)
    finalProp = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=fontProp), 
    			endParaRPr=fontProp)])
    rawChart.x_axis.txPr = finalProp
    rawChart.y_axis.txPr = finalProp
	
	#Adds data series to raw chart and sets cell for plotting the chart.
    xvalues = Reference(sheet, min_col=1, min_row=2, max_row=sheetCounter-1)
    yvalues = Reference(sheet, min_col=2, min_row=2, max_row=sheetCounter-1)
    series = Series(values = yvalues, xvalues = xvalues)
    rawChart.series.append(series)
    sheet.add_chart(rawChart, "E5")
    
    #Formats the raw chart with no gridlines.
    noGridLines = GraphicalProperties(ln=LineProperties(noFill=True))
    rawChart.x_axis.majorGridlines.spPr = noGridLines
    rawChart.y_axis.majorGridlines.spPr = noGridLines
    
    #Removes the legend from the raw chart.
    rawChart.legend = None
    
    #Sets the min and max values for the x-axis in the raw chart.
    rawChart.x_axis.scaling.min = minRT
    rawChart.x_axis.scaling.max = maxRT
    
    #Formats the color and width of the scatter plot line in the raw plot.
    series.graphicalProperties.line.solidFill = "0080FF"
    series.graphicalProperties.line.width = 100
    
    #Plots the normalized data via a scatter plot and sets titles for axes.
    normChart = ScatterChart()
    normChart.y_axis.title = "Normalized Absorbance"
    normChart.x_axis.title = "Retention Time (min)"
    
    #Sets the normalized chart title text font and size.
    normChart.x_axis.title.tx.rich.p[0].r.rPr = textProp
    normChart.y_axis.title.tx.rich.p[0].r.rPr = textProp
    
    #Sets the normalized axis text font and size.
    normChart.x_axis.txPr = finalProp
    normChart.y_axis.txPr = finalProp
	
	#Adds data series to normalized chart and sets cell for plotting the chart.
    xvalues = Reference(sheet, min_col=1, min_row=2, max_row=sheetCounter-1)
    yvalues = Reference(sheet, min_col=3, min_row=2, max_row=sheetCounter-1)
    series = Series(values = yvalues, xvalues = xvalues)
    normChart.series.append(series)
    sheet.add_chart(normChart, "N5")
    
    #Formats the normalized chart with no gridlines.
    normChart.x_axis.majorGridlines.spPr = noGridLines
    normChart.y_axis.majorGridlines.spPr = noGridLines
    
    #Removes the legend from the normalized chart.
    normChart.legend = None
    
    #Sets the min and max values for each axis in the normalized chart.
    normChart.y_axis.scaling.min = 0
    normChart.y_axis.scaling.max = 100
    normChart.x_axis.scaling.min = minRT
    normChart.x_axis.scaling.max = maxRT
    
    #Sets the y-axis major units on the normalized chart.
    normChart.y_axis.majorUnit = 100
    
    #Formats the color and width of the scatter plot line in the normalized plot.
    series.graphicalProperties.line.solidFill = "FF0000"
    series.graphicalProperties.line.width = 100
    
    #Saves the output Excel file.
    FileName = File[:-4] #Removes .CSV from File Name
    outFile.save("Plot for "+FileName+".xlsx")
    
#Lets user know that the program is finished.
finished = raw_input("ATM is finished! Press 'enter' to close the program.")